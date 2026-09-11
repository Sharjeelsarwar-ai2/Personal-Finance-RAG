import os
import re
import pandas as pd
import pdfplumber
# NEW (Works with both modern pypdf and PyPDF2):
try:
    from pypdf import PdfReader
except ImportError:
    from PyPDF2 import PdfReader
from typing import List, Dict, Tuple
import io


class DocumentProcessor:
    """Handles extraction of text and tables from various financial document formats"""

    SUPPORTED_FORMATS = ["pdf", "csv", "xlsx", "xls", "txt", "docx"]

    def __init__(self):
        self.processed_docs = []

    def process_file(self, uploaded_file) -> Dict:
        """
        Process a single uploaded file and extract content.
        Returns dict with filename, content, tables, and metadata.
        """
        filename = uploaded_file.name
        file_ext = filename.rsplit(".", 1)[-1].lower()
        file_size = uploaded_file.size / 1024  # KB

        result = {
            "filename": filename,
            "file_type": file_ext,
            "file_size": file_size,
            "text_content": "",
            "tables": [],
            "metadata": {},
            "raw_rows": []
        }

        try:
            if file_ext == "pdf":
                result = self._process_pdf(uploaded_file, result)
            elif file_ext == "csv":
                result = self._process_csv(uploaded_file, result)
            elif file_ext in ["xlsx", "xls"]:
                result = self._process_excel(uploaded_file, result)
            elif file_ext == "txt":
                result = self._process_txt(uploaded_file, result)
            elif file_ext == "docx":
                result = self._process_docx(uploaded_file, result)
            else:
                result["text_content"] = f"[Unsupported format: {file_ext}]"

            result["metadata"]["status"] = "success"

        except Exception as e:
            result["text_content"] = f"[Error processing {filename}: {str(e)}]"
            result["metadata"]["status"] = "error"
            result["metadata"]["error"] = str(e)

        self.processed_docs.append(result)
        return result

    def _process_pdf(self, uploaded_file, result: Dict) -> Dict:
        """Extract text and tables from PDF files"""
        text_parts = []
        tables_found = []

        # Reset file pointer
        uploaded_file.seek(0)
        file_bytes = uploaded_file.read()

        # --- Extract text using PyPDF2 ---
        uploaded_file.seek(0)
        reader = PdfReader(io.BytesIO(file_bytes))
        result["metadata"]["pages"] = len(reader.pages)

        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(f"[Page {i+1}]\n{page_text.strip()}")

        # --- Extract tables using pdfplumber ---
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_tables = page.extract_tables()
                    if page_tables:
                        for t_idx, table in enumerate(page_tables):
                            if table and len(table) > 1:
                                # Convert table to readable text
                                headers = table[0] if table[0] else [f"Col_{j}" for j in range(len(table[1]))]
                                headers = [str(h).strip() if h else f"Col_{j}" for j, h in enumerate(headers)]

                                table_text = f"\n[Table from Page {i+1}, Table {t_idx+1}]\n"
                                table_text += " | ".join(headers) + "\n"
                                table_text += "-" * 50 + "\n"

                                for row in table[1:]:
                                    if row:
                                        cleaned_row = [str(cell).strip() if cell else "" for cell in row]
                                        table_text += " | ".join(cleaned_row) + "\n"

                                tables_found.append(table_text)

                                # Also try to create DataFrame
                                try:
                                    df = pd.DataFrame(table[1:], columns=headers)
                                    result["raw_rows"].extend(df.to_dict("records"))
                                except Exception:
                                    pass
        except Exception:
            pass  # pdfplumber might fail on some PDFs; we still have PyPDF2 text

        result["text_content"] = "\n\n".join(text_parts)
        if tables_found:
            result["text_content"] += "\n\n--- EXTRACTED TABLES ---\n" + "\n".join(tables_found)
        result["tables"] = tables_found

        return result

    def _process_csv(self, uploaded_file, result: Dict) -> Dict:
        """Process CSV bank statements / financial data"""
        uploaded_file.seek(0)

        try:
            df = pd.read_csv(uploaded_file, encoding="utf-8")
        except UnicodeDecodeError:
            uploaded_file.seek(0)
            df = pd.read_csv(uploaded_file, encoding="latin-1")

        result = self._dataframe_to_result(df, result)
        return result

    def _process_excel(self, uploaded_file, result: Dict) -> Dict:
        """Process Excel files"""
        uploaded_file.seek(0)

        xls = pd.ExcelFile(uploaded_file)
        all_text = []
        all_rows = []

        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            if df.empty:
                continue

            sheet_text = f"\n[Sheet: {sheet_name}]\n"
            sheet_text += self._dataframe_to_text(df)
            all_text.append(sheet_text)
            all_rows.extend(df.to_dict("records"))

        result["text_content"] = "\n\n".join(all_text)
        result["raw_rows"] = all_rows
        result["metadata"]["sheets"] = xls.sheet_names
        result["metadata"]["total_rows"] = len(all_rows)

        return result

    def _process_txt(self, uploaded_file, result: Dict) -> Dict:
        """Process plain text files"""
        uploaded_file.seek(0)
        content = uploaded_file.read().decode("utf-8", errors="ignore")
        result["text_content"] = content.strip()
        result["metadata"]["char_count"] = len(content)
        return result

    def _process_docx(self, uploaded_file, result: Dict) -> Dict:
        """Process Word documents"""
        from docx import Document

        uploaded_file.seek(0)
        doc = Document(io.BytesIO(uploaded_file.read()))

        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())

        # Extract tables from docx
        for t_idx, table in enumerate(doc.tables):
            table_text = f"\n[Document Table {t_idx+1}]\n"
            rows_data = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_cells)
                table_text += " | ".join(row_cells) + "\n"
            paragraphs.append(table_text)

        result["text_content"] = "\n\n".join(paragraphs)
        result["metadata"]["paragraphs"] = len(paragraphs)

        return result

    def _dataframe_to_result(self, df: pd.DataFrame, result: Dict) -> Dict:
        """Convert a DataFrame to text representation for RAG"""
        result["text_content"] = self._dataframe_to_text(df)
        result["raw_rows"] = df.to_dict("records")
        result["metadata"]["columns"] = list(df.columns)
        result["metadata"]["total_rows"] = len(df)

        # Basic financial stats
        numeric_cols = df.select_dtypes(include=["number"]).columns.tolist()
        if numeric_cols:
            stats = {}
            for col in numeric_cols:
                stats[col] = {
                    "sum": float(df[col].sum()),
                    "mean": float(df[col].mean()),
                    "min": float(df[col].min()),
                    "max": float(df[col].max())
                }
            result["metadata"]["numeric_stats"] = stats

        return result

    def _dataframe_to_text(self, df: pd.DataFrame) -> str:
        """Convert DataFrame to structured text for embedding"""
        text_parts = []

        # Column info
        text_parts.append(f"Columns: {', '.join(df.columns.tolist())}")
        text_parts.append(f"Total rows: {len(df)}")

        # Summary statistics for numeric columns
        numeric_cols = df.select_dtypes(include=["number"]).columns.tolist()
        if numeric_cols:
            text_parts.append("\n--- SUMMARY STATISTICS ---")
            for col in numeric_cols:
                text_parts.append(
                    f"{col}: Total={df[col].sum():.2f}, "
                    f"Average={df[col].mean():.2f}, "
                    f"Min={df[col].min():.2f}, "
                    f"Max={df[col].max():.2f}"
                )

        # Row-by-row data (convert each row to natural language)
        text_parts.append("\n--- TRANSACTION DATA ---")

        for idx, row in df.iterrows():
            row_text = f"Row {idx+1}: "
            row_parts = []
            for col in df.columns:
                val = row[col]
                if pd.notna(val):
                    row_parts.append(f"{col}={val}")
            row_text += ", ".join(row_parts)
            text_parts.append(row_text)

            # Safety: limit to 500 rows for very large files
            if idx >= 499:
                text_parts.append(f"[... truncated, showing 500 of {len(df)} rows ...]")
                break

        return "\n".join(text_parts)

    def get_all_text(self) -> str:
        """Get combined text from all processed documents"""
        all_texts = []
        for doc in self.processed_docs:
            header = f"\n{'='*60}\nDOCUMENT: {doc['filename']} ({doc['file_type'].upper()})\n{'='*60}\n"
            all_texts.append(header + doc["text_content"])
        return "\n\n".join(all_texts)

    def get_doc_summaries(self) -> List[Dict]:
        """Get summary info for all processed docs"""
        summaries = []
        for doc in self.processed_docs:
            summaries.append({
                "filename": doc["filename"],
                "file_type": doc["file_type"],
                "file_size": doc["file_size"],
                "content_length": len(doc["text_content"]),
                "tables_found": len(doc["tables"]),
                "status": doc["metadata"].get("status", "unknown")
            })
        return summaries

    def reset(self):
        """Clear all processed documents"""
        self.processed_docs = []
